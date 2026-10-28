"""把工具结果（Markdown）转成 Word 文档。用户点「下载 Word」时才现转，不预先生成、不落盘。

支持结果里实际会出现的语法：标题、段落、有序/无序列表（含缩进）、加粗/斜体/行内代码、表格、引用、
分隔线、代码块、勾选框。中文字体统一设好，Word 打开不会出现字体错乱。

导出方式（mode）：
  full      完整结果
  resume    只导出「完整优化版简历」那一节：简历修改类工具用，拿到的是能直接投的简历
  practice  每道题下面留作答框：面试题类工具用，打印出来就能练
"""

import io
import re
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

CN_FONT = "Microsoft YaHei"     # Windows 自带；Mac 上的 Word 会自动换成苹方
EN_FONT = "Calibri"
GRAY = RGBColor(0x6B, 0x72, 0x80)
INK = RGBColor(0x11, 0x18, 0x27)
WATER = RGBColor(0x00, 0x77, 0xB6)

# 简历修改结果里「能直接投的那份简历」所在的标题
RESUME_HEADING = re.compile(r"(完整|最终)?(优化|修改|改写)(后|版)的?简历|简历(全文|正文|终稿)")
# 面试题结果里「一道题」的标题：Q1. / P2. / 问题 3 / 1.
QUESTION_HEADING = re.compile(r"^([QP]\d+|问题\s*\d+|第\s*\d+\s*题|\d+[.、])")

_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")
_LIST = re.compile(r"^(\s*)([-*+•]|\d+[.、)])\s+(.*)$")
_TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _font(target, name: str = CN_FONT, size: Optional[float] = None, color: Optional[RGBColor] = None,
          bold: Optional[bool] = None) -> None:
    """同时设西文和中文字体：只设 font.name 的话，中文还是 Word 的默认宋体。"""
    f = target.font
    f.name = EN_FONT if name == CN_FONT else name
    rpr = target.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)
    if size:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold


def _setup(doc: Document) -> None:
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    _font(doc.styles["Normal"], size=10.5, color=INK)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.35
    for lvl, size in ((1, 16), (2, 13.5), (3, 11.5), (4, 10.5)):
        st = doc.styles[f"Heading {lvl}"]
        _font(st, size=size, color=WATER if lvl <= 2 else INK, bold=True)
        st.paragraph_format.space_before = Pt(12 if lvl <= 2 else 8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def _inline(par, text: str, bold: bool = False, color: Optional[RGBColor] = None) -> None:
    text = text.replace("[ ] ", "☐ ").replace("[x] ", "☑ ").replace("[X] ", "☑ ")
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i, code = bold, False, False
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            part = part[1:part.index("](")]
        run = par.add_run(part)
        run.bold, run.italic = b or None, i or None
        if code:
            _font(run, "Consolas", color=RGBColor(0xC0, 0x26, 0x5A))
        elif color is not None:
            run.font.color.rgb = color


def _cells(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _table(doc: Document, rows: List[List[str]]) -> None:
    width = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=width)
    t.style = "Table Grid"
    for ri, r in enumerate(rows):
        for ci in range(width):
            cell = t.cell(ri, ci)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            _inline(par, r[ci] if ci < len(r) else "", bold=(ri == 0))
            for run in par.runs:
                run.font.size = Pt(9.5)
    doc.add_paragraph()


def _answer_box(doc: Document) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    p = cell.paragraphs[0]
    r = p.add_run("我的回答：")
    _font(r, size=9.5, color=GRAY)
    for _ in range(5):
        cell.add_paragraph()
    doc.add_paragraph()


def resume_section(md: str) -> Optional[str]:
    """取出「完整优化版简历」那一节，到同级或更高级标题为止。没有这一节返回 None。"""
    lines = md.splitlines()
    for i, l in enumerate(lines):
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m and RESUME_HEADING.search(m.group(2)):
            lvl, body = len(m.group(1)), []
            for x in lines[i + 1:]:
                n = re.match(r"^(#{1,4})\s+", x)
                if n and len(n.group(1)) <= lvl:
                    break
                body.append(x)
            text = "\n".join(body).strip()
            return text or None
    return None


def build(md: str, title: str, meta: str, mode: str = "full") -> bytes:
    doc = Document()
    _setup(doc)
    if mode == "resume":
        md = resume_section(md) or md
    else:
        h = doc.add_paragraph()
        _font(h.add_run(title), size=18, color=INK, bold=True)
        if meta:
            m = doc.add_paragraph()
            r = m.add_run(meta)
            _font(r, size=9, color=GRAY)

    lines = md.splitlines()
    i, pending_answer = 0, False

    def flush_answer():
        nonlocal pending_answer
        if pending_answer:
            _answer_box(doc)
            pending_answer = False

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                _font(p.add_run(lines[i]), "Consolas", size=9)
                i += 1
            i += 1
            continue
        hm = re.match(r"^(#{1,6})\s+(.*)", s)
        if hm:
            flush_answer()
            lvl = min(len(hm.group(1)), 4)
            text = hm.group(2).strip()
            p = doc.add_heading(level=lvl)
            _inline(p, text)
            if mode == "practice" and lvl >= 2 and QUESTION_HEADING.match(text):
                pending_answer = True
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1].strip()):
            rows = [_cells(s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            _table(doc, rows)
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            i += 1
            continue   # 分隔线：Word 里靠标题分段就够了，画线反而乱
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            _inline(p, s.lstrip("> ").strip(), color=GRAY)
            i += 1
            continue
        lm = _LIST.match(line)
        if lm:
            depth = len(lm.group(1).replace("\t", "  ")) // 2
            marker, text = lm.group(2), lm.group(3)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(2)
            bullet = marker if marker[0].isdigit() else ("•" if depth == 0 else "◦")
            p.add_run(bullet + " ")
            _inline(p, text)
            i += 1
            continue
        if not s:
            i += 1
            continue
        p = doc.add_paragraph()
        _inline(p, s)
        i += 1
    flush_answer()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
