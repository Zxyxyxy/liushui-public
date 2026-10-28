"""文件上传：校验 → 存盘 → 提取文字 → 登记。

接口：
    POST   /api/upload        multipart，字段 file；需登录
    GET    /api/upload/<id>/text  取自己文件的提取文字（工具页填进表单用）
    DELETE /api/upload/<id>   删除自己的文件（磁盘原件 + 提取的文字）

安全取舍：
- 类型看内容（文件头魔数），不看扩展名；docx 还要确认是合法 zip 且含 word/document.xml。
- 磁盘文件名由服务端生成（uuid），用户给的文件名只清洗后用于展示，永远不参与路径拼接。
- 文件按用户分目录；接口只返回 id / 展示名 / 字数，不返回磁盘路径。
- 解析放在子进程里跑并设超时：畸形 PDF 可能让解析器卡死或吃光内存，不能拖垮主服务。
- 提取的文字当作「用户内容」交给模型，不当指令——这一点由 agent 的提示词和脱敏层保证。
- 保留 UPLOAD_TTL_DAYS 天（默认 7）。过期清理在服务启动和每次上传时顺带做。
"""
import os
import re
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from flask import Blueprint, jsonify, request

from auth import auth_required_response, current_user
from db import (add_upload, count_live_uploads, get_upload, pop_expired_uploads,
                pop_upload)

bp = Blueprint("uploads", __name__)

UPLOAD_DIR = Path(__file__).parent / "uploads_data"
MAX_BYTES = 10 * 1024 * 1024          # 单文件 10MB
MAX_TEXT_CHARS = 30000                # 提取文字上限；再长的简历/资料基本不是一份简历
MAX_LIVE_PER_USER = 20                # 每人同时保留的文件数
UPLOAD_TTL_DAYS = int(os.environ.get("UPLOAD_TTL_DAYS", "7"))
PARSE_TIMEOUT = 20                    # 秒
RATE_LIMIT = 10                       # 每分钟每人最多上传次数
RATE_WINDOW = 60

_hits: Dict[str, List[float]] = {}     # 进程内限流，多实例时需换共享存储（同登录限流）


def _rate_limited(user_id: str) -> bool:
    now = time.time()
    recent = [t for t in _hits.get(user_id, []) if now - t < RATE_WINDOW]
    if len(recent) >= RATE_LIMIT:
        _hits[user_id] = recent
        return True
    recent.append(now)
    _hits[user_id] = recent
    return False


def sniff_kind(head: bytes, full_path: Path) -> Optional[str]:
    """按内容判断类型，只认 pdf / docx。"""
    if head.startswith(b"%PDF-"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(full_path) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def clean_name(raw: str) -> str:
    """展示用文件名：去掉路径成分和控制字符，限长。不参与任何路径操作。"""
    name = (raw or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f<>\"|?*]", "", name).strip().strip(".")
    return (name or "未命名文件")[:80]


# ── 文字提取（在子进程里执行）──────────────────────────────────

_WORKER = r'''
import sys, json
path, kind = sys.argv[1], sys.argv[2]
out = []
if kind == "pdf":
    from pypdf import PdfReader
    r = PdfReader(path)
    if r.is_encrypted:
        print(json.dumps({"error": "encrypted"})); sys.exit(0)
    for page in r.pages[:60]:
        out.append(page.extract_text() or "")
else:
    import docx
    d = docx.Document(path)
    for p in d.paragraphs:
        out.append(p.text)
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
print(json.dumps({"text": "\n".join(out)}, ensure_ascii=False))
'''


def extract_text(path: Path, kind: str) -> Tuple[Optional[str], Optional[str]]:
    """返回 (文字, 错误文案)。错误文案可直接给用户看。"""
    try:
        # 不能用 -I：它会连用户目录（pip --user 装的包）一起屏蔽，找不到 pypdf。
        # 改用 -E 忽略环境变量，并把工作目录设成一个空目录——
        # -c 会把当前目录加进模块搜索路径，空目录保证没人能借此塞同名模块进来。
        with tempfile.TemporaryDirectory() as empty:
            proc = subprocess.run(
                [sys.executable, "-E", "-c", _WORKER, str(path), kind],
                capture_output=True, timeout=PARSE_TIMEOUT, text=True, cwd=empty,
            )
    except subprocess.TimeoutExpired:
        return None, "文件解析超时，请换一个文件或转成文字版再试"
    if proc.returncode != 0:
        if "No module named" in (proc.stderr or ""):
            return None, "服务器还没装好文件解析组件，请联系管理员"
        return None, "这个文件打不开，可能已损坏，请换一个"
    try:
        import json
        data = json.loads(proc.stdout.strip().splitlines()[-1])
    except Exception:
        return None, "这个文件打不开，可能已损坏，请换一个"
    if data.get("error") == "encrypted":
        return None, "文件有密码保护，请去掉密码后再上传"
    text = re.sub(r"\n{3,}", "\n\n", (data.get("text") or "")).strip()
    if len(text) < 20:
        return None, "没能从文件里读出文字。可能是扫描件或图片版 PDF，请改传文字版，或直接粘贴内容"
    return text, None


# ── 清理 ────────────────────────────────────────────────────

def _file_path(user_id: str, stored_name: str) -> Path:
    # stored_name 来自服务端生成；这里再校验一次，防御库里被写入异常值
    if not re.fullmatch(r"[0-9a-f-]{36}\.(pdf|docx)", stored_name):
        raise ValueError("bad stored_name")
    if not re.fullmatch(r"[0-9a-f-]{36}", user_id):
        raise ValueError("bad user_id")
    return UPLOAD_DIR / user_id / stored_name


def _remove_file(row: dict) -> None:
    try:
        _file_path(row["user_id"], row["stored_name"]).unlink(missing_ok=True)
    except (ValueError, OSError):
        pass


def purge_expired() -> int:
    rows = pop_expired_uploads()
    for r in rows:
        _remove_file(r)
    return len(rows)


# ── 对外：按 id 取文字（对话 / 工具运行用）──────────────────────

def load_upload_text(upload_id: str, user_id: str) -> Optional[dict]:
    """只能取到自己的、未过期的文件。返回 {id,name,text,truncated} 或 None。"""
    row = get_upload(upload_id, user_id)
    if row is None:
        return None
    return {"id": row["id"], "name": row["name"], "text": row["text"],
            "truncated": bool(row["truncated"])}


# ── 接口 ────────────────────────────────────────────────────

def _public(row: dict) -> dict:
    return {"id": row["id"], "name": row["name"], "kind": row["kind"],
            "size": row["size"], "chars": row["chars"],
            "truncated": bool(row["truncated"]), "expires_at": row["expires_at"]}


@bp.route("/api/upload", methods=["POST"])
def upload():
    user = current_user()
    if not user:
        return auth_required_response()
    uid = user["id"]

    if _rate_limited(uid):
        return jsonify({"success": False, "error": "上传太频繁了，请稍后再试"}), 429
    f = request.files.get("file")
    if f is None or not f.filename:
        return jsonify({"success": False, "error": "没有收到文件"}), 400

    purge_expired()
    if count_live_uploads(uid) >= MAX_LIVE_PER_USER:
        return jsonify({"success": False,
                        "error": f"最多同时保留 {MAX_LIVE_PER_USER} 个文件，请先删除一些"}), 400

    user_dir = UPLOAD_DIR / uid
    user_dir.mkdir(parents=True, exist_ok=True)
    tmp = user_dir / f"{uuid.uuid4()}.part"
    try:
        # 分块写入并计数：不信 Content-Length，超限立刻停
        size = 0
        with open(tmp, "wb") as out:
            while True:
                chunk = f.stream.read(64 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_BYTES:
                    return jsonify({"success": False, "error": "文件超过 10MB，请压缩或精简后再传"}), 413
                out.write(chunk)
        if size == 0:
            return jsonify({"success": False, "error": "文件是空的"}), 400

        with open(tmp, "rb") as fh:
            head = fh.read(8)
        kind = sniff_kind(head, tmp)
        if kind is None:
            return jsonify({"success": False,
                            "error": "目前只支持 PDF 和 Word（.docx）文件"}), 415

        text, err = extract_text(tmp, kind)
        if err:
            return jsonify({"success": False, "error": err}), 422

        truncated = len(text) > MAX_TEXT_CHARS
        if truncated:
            text = text[:MAX_TEXT_CHARS]

        stored = f"{uuid.uuid4()}.{kind}"
        os.replace(tmp, user_dir / stored)
        row = add_upload(uid, clean_name(f.filename), stored, kind, size, text,
                         truncated, UPLOAD_TTL_DAYS)
        return jsonify({"success": True, "file": _public(row)})
    finally:
        tmp.unlink(missing_ok=True)


@bp.route("/api/upload/<upload_id>", methods=["DELETE"])
def delete_upload(upload_id):
    user = current_user()
    if not user:
        return auth_required_response()
    row = pop_upload(upload_id, user["id"])
    if row is None:
        return jsonify({"success": False, "error": "文件不存在"}), 404
    _remove_file(row)
    return jsonify({"success": True})


@bp.route("/api/upload/<upload_id>/text", methods=["GET"])
def upload_text(upload_id):
    user = current_user()
    if not user:
        return auth_required_response()
    got = load_upload_text(upload_id, user["id"])
    if got is None:
        return jsonify({"success": False, "error": "文件不存在或已过期"}), 404
    return jsonify({"success": True, "text": got["text"], "truncated": got["truncated"]})
